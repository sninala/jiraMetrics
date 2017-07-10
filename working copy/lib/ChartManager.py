from Constants import Constants
from openpyxl.chart.trendline import Trendline
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.colors import ColorChoice
from openpyxl.chart import BarChart, LineChart, Series, Reference


class ChartManager(object):

    def __init__(self, config, data_sheet, charts_sheet):
        self.config = config
        self.data_sheet = data_sheet
        self.charts_sheet = charts_sheet

    def draw_barchart(self, chart_properties):
        print "Creating Bar chart for {}".format(chart_properties['title'])
        chart = BarChart()
        chart.height = 12
        chart.width = 30
        chart.style = 10
        chart.title = chart_properties['title']

        data = Reference(
            self.data_sheet, min_col=chart_properties['data_min_column'],
            min_row=chart_properties['data_min_row'],
            max_row=chart_properties['data_max_row'],
            max_col=chart_properties['data_max_column'])
        cats = Reference(
            self.data_sheet,
            min_col=chart_properties['cats_min_column'],
            min_row=chart_properties['cats_min_row'],
            max_row=chart_properties['cats_max_row'])
        chart.add_data(data, titles_from_data=True)
        if "stacked" in chart_properties:
            chart.grouping = "stacked"
        if "y_axis_min_value" in chart_properties:
            chart.y_axis.scaling.min = chart_properties["y_axis_min_value"]
        if "y_axis_max_value" in chart_properties:
            chart.y_axis.scaling.max = chart_properties["y_axis_max_value"]
        if chart_properties['logarithmic_y_axis']:
            chart.y_axis.scaling.logBase = 10
            chart.y_axis.minorTickMark = "out"
        chart.set_categories(cats)
        chart.shape = 4
        projects = chart_properties['projects']
        series = chart.series
        for index, current_series in enumerate(series):
            if projects:
                project = projects[index]
                project_color = self.config.get('PROJECT_COLOR', project)
                current_series.graphicalProperties.solidFill = project_color
            else:
                current_series.graphicalProperties.solidFill = "1381BD"

            if chart_properties['trendline']:
                current_series.trendline = Trendline()
                current_series.trendline.trendlineType = 'linear'
                line_properties = LineProperties(w=22250, solidFill=ColorChoice(prstClr='black'))
                current_series.trendline.graphicalProperties = GraphicalProperties(ln=line_properties)

        if chart_properties['data_labels']:
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
        self.charts_sheet.add_chart(chart, chart_properties['cell'])

    def draw_linechart(self, chart_properties):
        print "Creating line chart for {}".format(chart_properties['title'])
        chart = LineChart()
        chart.height = 12
        chart.width = 30
        chart.title = chart_properties['title']
        chart.style = 12
        chart.x_axis.tickLblPos = "low"
        if chart_properties['logarithmic_y_axis']:
            chart.y_axis.scaling.logBase = 10
            chart.y_axis.minorTickMark = "out"
        if "y_axis_min_value" in chart_properties:
            chart.y_axis.scaling.min = chart_properties["y_axis_min_value"]
        if "y_axis_max_value" in chart_properties:
            chart.y_axis.scaling.max = chart_properties["y_axis_max_value"]
        data = Reference(
            self.data_sheet, min_col=chart_properties['data_min_column'],
            min_row=chart_properties['data_min_row'],
            max_row=chart_properties['data_max_row'],
            max_col=chart_properties['data_max_column'])
        cats = Reference(
            self.data_sheet,
            min_col=chart_properties['cats_min_column'],
            min_row=chart_properties['cats_min_row'],
            max_row=chart_properties['cats_max_row'])

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        projects = chart_properties['projects']
        statistics = chart_properties['statistics']
        series = chart.series
        for index, current_series in enumerate(series):
            if projects:
                project = projects[index]
                project_color = self.config.get('PROJECT_COLOR', project)
                marker_symbol = self.config.get('PROJECT_MARKER_SYMBOL', project)
                current_series.marker.symbol = marker_symbol
                current_series.marker.size = 7
                if marker_symbol in ("triangle", "diamond", "circle"):
                    current_series.marker.graphicalProperties.solidFill = project_color # Marker filling
                current_series.marker.graphicalProperties.line.solidFill = project_color
                current_series.graphicalProperties.line.solidFill = project_color
            elif statistics:
                statistic = statistics[index]
                properties = self.get_chart_properties_for(statistic)
                current_series.marker.symbol = properties["MARKER_SYMBOL"]
                current_series.marker.size = 7
                if properties["MARKER_SYMBOL"] in ("triangle", "diamond", "circle"):
                    current_series.marker.graphicalProperties.solidFill = properties["COLOR"]  # Marker filling
                current_series.marker.graphicalProperties.line.solidFill = properties["COLOR"]
                current_series.graphicalProperties.line.solidFill = properties["COLOR"]
            else:
                current_series.marker.symbol = "diamond"
                current_series.marker.graphicalProperties.solidFill = "1381BD"  # Marker filling
                current_series.marker.graphicalProperties.line.solidFill = "1381BD"
                current_series.graphicalProperties.line.solidFill = "1381BD"
            current_series.graphicalProperties.line.width = 28568
            if chart_properties['trendline']:
                current_series.trendline = Trendline()
                current_series.trendline.trendlineType = 'linear'
                line_properties = LineProperties(w=22250, solidFill=ColorChoice(prstClr='black'))
                current_series.trendline.graphicalProperties = GraphicalProperties(ln=line_properties)
        if chart_properties['data_labels']:
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
            chart.dataLabels.position = "b"

        self.charts_sheet.add_chart(chart, chart_properties['cell'])

    def get_chart_properties_for(self, metric):
        chart_props = Constants.METRIC_PROPERTIES
        return chart_props[metric]


